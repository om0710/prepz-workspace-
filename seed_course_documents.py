import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sqlite3
from datetime import datetime

os.makedirs("uploads", exist_ok=True)

SEED_DOCS = [
    {
        "filename": "CSET204_Probability_and_Statistics_EndSem_2025.docx",
        "subject": "Probability & Statistics",
        "semester": "Semester 3",
        "file_type": "PYQ",
        "exam_type": "End-Sem",
        "uploader_name": "Om Bansal (Lead Contributor)",
        "uploader_email": "ombansal221@gmail.com",
        "title": "Bennett University - End Semester Examination (2025)\nCourse Code: CSET204 | Probability and Statistics",
        "sections": [
            ("Section A: Short Conceptual Questions (2 Marks Each)", [
                "1. State Bayes' Theorem and write its mathematical formulation for mutually exclusive events.",
                "2. Define the Probability Density Function (PDF) of a continuous random variable and state its two fundamental properties.",
                "3. If a fair die is rolled 5 times, find the probability of obtaining exactly 2 sixes using Binomial distribution.",
                "4. Differentiate between Type I and Type II errors in Statistical Hypothesis Testing.",
                "5. What is the variance of a Poisson distribution with parameter lambda = 4?"
            ]),
            ("Section B: Analytical & Numerical Problems (7 Marks Each)", [
                "1. A manufacturing process produces 5% defective microchips. In a random sample of 200 microchips, use Poisson approximation to find the probability that (a) exactly 3 are defective, (b) at least 2 are defective.",
                "2. The joint probability mass function of (X, Y) is given by P(X=x, Y=y) = c*(2x + y) for x = 0,1,2 and y = 0,1,2,3. Find (a) constant c, (b) Marginal distributions of X and Y, (c) Covariance Cov(X, Y).",
                "3. A random variable X has normal distribution with mean 50 and standard deviation 10. Find P(X > 65), P(35 < X < 60), and the value k such that P(X < k) = 0.95.",
                "4. A sample of 100 students had an average score of 72 with standard deviation 8. Test the hypothesis that the population mean is 70 at 5% level of significance."
            ]),
            ("Section C: Comprehensive Derivations & Applications (14 Marks Each)", [
                "1. (a) Derive the Moment Generating Function (MGF) of an Exponential distribution and use it to find the Mean and Variance. (b) Explain Central Limit Theorem (CLT) and demonstrate its application in large-sample confidence intervals."
            ])
        ]
    },
    {
        "filename": "CSET204_Probability_and_Statistics_Complete_Notes.docx",
        "subject": "Probability & Statistics",
        "semester": "Semester 3",
        "file_type": "Notes",
        "exam_type": "Other",
        "uploader_name": "Aryan Sharma (CSE '25)",
        "uploader_email": "aryan.sharma@bennett.edu.in",
        "title": "Comprehensive Study Notes: Probability & Engineering Statistics\nBennett University School of Computer Science Engineering",
        "sections": [
            ("Unit 1: Probability Foundations & Conditional Probability", [
                "Axioms of Probability: Non-negativity, Normalization P(S) = 1, Additivity for disjoint events.",
                "Multiplication Rule: P(A ∩ B) = P(A) * P(B | A).",
                "Total Probability Law: P(B) = Sum[P(A_i) * P(B | A_i)] over all partitions A_i.",
                "Bayes Rule: P(A_k | B) = [P(A_k) * P(B | A_k)] / Sum[P(A_i) * P(B | A_i)]."
            ]),
            ("Unit 2: Discrete & Continuous Random Variables", [
                "Binomial: P(X=k) = nCk * p^k * (1-p)^(n-k). Mean = np, Var = np(1-p).",
                "Poisson: P(X=k) = (e^(-lambda) * lambda^k) / k!. Mean = lambda, Var = lambda.",
                "Exponential: f(x) = lambda * e^(-lambda*x) for x >= 0. Mean = 1/lambda, Memoryless property.",
                "Normal Distribution: Z = (X - mu) / sigma. Standard normal distribution Z ~ N(0, 1)."
            ]),
            ("Unit 3: Hypothesis Testing & Regression Analysis", [
                "Null Hypothesis (H0) vs Alternative Hypothesis (H1).",
                "Z-Test for large samples (n >= 30), t-Test for small samples (n < 30) with unknown variance.",
                "Linear Regression: y = beta_0 + beta_1 * x + epsilon, where beta_1 = Cov(X,Y) / Var(X)."
            ])
        ]
    },
    {
        "filename": "CSET101_Computational_Thinking_Python_PYQ.docx",
        "subject": "Computational Thinking & Programming",
        "semester": "Semester 1",
        "file_type": "PYQ",
        "exam_type": "End-Sem",
        "uploader_name": "Priya Patel (AI/DS '26)",
        "uploader_email": "priya.patel@bennett.edu.in",
        "title": "Bennett University - End Semester Examination\nCourse Code: CSET101 | Computational Thinking & Python Programming",
        "sections": [
            ("Section A: Short Python Concept Questions (2 Marks Each)", [
                "1. Explain the difference between mutable (lists, dicts) and immutable (tuples, strings) data types in Python with examples.",
                "2. What is list comprehension? Write a one-line Python list comprehension to filter out odd numbers from 1 to 50.",
                "3. Explain how dictionary hashing works and why dictionary lookups are O(1) average time complexity.",
                "4. Differentiate between '==' (equality of value) and 'is' (identity/memory address equality)."
            ]),
            ("Section B: Python Coding & Algorithms (7 Marks Each)", [
                "1. Write a Python function `frequency_analyzer(text)` that reads a paragraph and returns a dictionary of words sorted by highest frequency.",
                "2. Implement binary search in Python using both iterative and recursive approaches. State time and space complexities.",
                "3. Write a Python program to read a CSV file of student marks and calculate mean, median, and grade distribution."
            ])
        ]
    },
    {
        "filename": "MATH101_Engineering_Calculus_MidSem_PYQ.docx",
        "subject": "Engineering Calculus",
        "semester": "Semester 1",
        "file_type": "PYQ",
        "exam_type": "Mid-Sem",
        "uploader_name": "Rohan Mehta (ECE '25)",
        "uploader_email": "rohan.mehta@bennett.edu.in",
        "title": "Bennett University - Mid Semester Examination\nCourse Code: MATH101 | Engineering Calculus",
        "sections": [
            ("Section A: Calculus Fundamentals (3 Marks Each)", [
                "1. State Rolle's Theorem and verify it for f(x) = x^2 - 4x + 3 in the interval [1, 3].",
                "2. Evaluate limit x->0 of (sin(x) - x) / x^3 using L'Hopital's Rule.",
                "3. State Euler's Theorem for homogeneous functions of degree n in x and y."
            ]),
            ("Section B: Multi-Variable Calculus & Integrals (8 Marks Each)", [
                "1. Find the maxima and minima of f(x, y) = x^3 + y^3 - 3axy using second derivative test.",
                "2. Evaluate double integral of (x^2 + y^2) dx dy over the circular region x^2 + y^2 <= R^2 by converting to polar coordinates.",
                "3. Expand f(x, y) = e^(x+y) in powers of x and y up to second-degree terms using Taylor's Series."
            ])
        ]
    },
    {
        "filename": "CSET201_OOPS_Java_Exam_Questions_and_Notes.docx",
        "subject": "OOPS using Java",
        "semester": "Semester 3",
        "file_type": "Notes",
        "exam_type": "Mid-Sem",
        "uploader_name": "Sneha Gupta (CSE '26)",
        "uploader_email": "sneha.gupta@bennett.edu.in",
        "title": "Object Oriented Programming in Java - Core Exam Revision Notes\nBennett University Department of Computer Science",
        "sections": [
            ("Pillar 1: Encapsulation & Abstraction", [
                "Encapsulation: Bundling data (fields) and methods that operate on data within a single class using private access modifiers and public getters/setters.",
                "Abstract Classes vs Interfaces: Abstract class can have state, constructors, and concrete methods. Interfaces define pure contracts and support multiple inheritance."
            ]),
            ("Pillar 2: Inheritance & Polymorphism", [
                "Method Overloading (Compile-time Polymorphism): Same method name, different parameter signature.",
                "Method Overriding (Runtime Polymorphism): Child class provides specific implementation for method in parent class using @Override annotation.",
                "Dynamic Method Dispatch: Runtime resolution of overridden method calls via superclass reference variable."
            ]),
            ("Java Memory Management & Exceptions", [
                "Heap (Object instances) vs Stack (Primitive local variables and method call stack frames).",
                "Try-Catch-Finally block execution order and custom Checked vs Unchecked exceptions."
            ])
        ]
    },
    {
        "filename": "CSET210_Operating_Systems_Process_and_Deadlocks_Notes.docx",
        "subject": "Operating Systems",
        "semester": "Semester 4",
        "file_type": "Notes",
        "exam_type": "Mid-Sem",
        "uploader_name": "Om Bansal (Lead Contributor)",
        "uploader_email": "ombansal221@gmail.com",
        "title": "Operating Systems - Process Scheduling, Concurrency & Deadlocks\nBennett University Engineering Revision",
        "sections": [
            ("Process Scheduling Algorithms", [
                "FCFS (First Come First Serve): Non-preemptive, suffers from Convoy Effect.",
                "SJF / SRTF (Shortest Job First / Shortest Remaining Time First): Provably optimal average waiting time, risk of starvation.",
                "Round Robin (RR): Preemptive using time quantum q. Performance depends heavily on time quantum size."
            ]),
            ("Process Synchronization & Semaphores", [
                "Critical Section Problem: Mutual Exclusion, Progress, Bounded Waiting.",
                "Counting Semaphores: wait(S) decrements, signal(S) increments. If S < 0, magnitude indicates number of blocked processes.",
                "Peterson's Solution and Test-and-Set hardware atomic instructions."
            ]),
            ("Deadlocks & Banker's Algorithm", [
                "4 Coffman Conditions: Mutual Exclusion, Hold & Wait, No Preemption, Circular Wait.",
                "Banker's Algorithm: Resource-allocation and safety evaluation using Need = Max - Allocation."
            ])
        ]
    }
]

def generate_docs():
    conn = sqlite3.connect("chatbot.db")
    cur = conn.cursor()

    for item in SEED_DOCS:
        file_path = os.path.join("uploads", item["filename"])
        doc = docx.Document()
        
        # Header title
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(item["title"])
        title_run.font.size = Pt(16)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(255, 107, 0)
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        doc.add_paragraph("─" * 60)
        
        meta_p = doc.add_paragraph()
        meta_p.add_run(f"Subject: {item['subject']} | Semester: {item['semester']} | Type: {item['file_type']} ({item['exam_type']})\n").bold = True
        meta_p.add_run(f"Uploaded by: {item['uploader_name']} ({item['uploader_email']})\nVerified by: Bennett University Peer Academic Review")
        
        doc.add_paragraph("─" * 60)

        for sec_title, bullets in item["sections"]:
            h = doc.add_heading(sec_title, level=2)
            for b in bullets:
                p = doc.add_paragraph(style='List Bullet')
                p.add_run(b)

        doc.save(file_path)
        file_size = os.path.getsize(file_path)
        print(f"Generated: {file_path} ({file_size} bytes)")

        cur.execute("""
            INSERT OR REPLACE INTO user_uploads (filename, user_email, user_name, uploaded_at, file_path, size_bytes, subject, semester, file_type, exam_type, is_private)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0)
        """, (
            item["filename"],
            item["uploader_email"],
            item["uploader_name"],
            datetime.now().isoformat(),
            file_path,
            file_size,
            item["subject"],
            item["semester"],
            item["file_type"],
            item["exam_type"]
        ))

    conn.commit()
    conn.close()
    print("Database seeding completed successfully!")

if __name__ == "__main__":
    generate_docs()
